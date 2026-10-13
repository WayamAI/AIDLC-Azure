"""Uploaded documents are parsed by type instead of decoded as raw bytes."""
import io

import pytest
from docx import Document
from fastapi import HTTPException

from app.routes.impact import extract_document_text


def _docx_bytes(*paragraphs: str) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cell-a"
    table.rows[0].cells[1].text = "cell-b"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_docx_text_is_extracted():
    text = extract_document_text("Spec.DOCX", _docx_bytes("User can reset password", "Link expires in 1h"))

    assert "User can reset password" in text
    assert "Link expires in 1h" in text
    assert "cell-a | cell-b" in text
    assert "PK" not in text


def test_plain_text_is_decoded():
    assert extract_document_text("notes.md", "# Títle\nbody".encode("utf-8")) == "# Títle\nbody"
    assert extract_document_text("notes.txt", "héllo".encode("utf-16")) == "héllo"


@pytest.mark.parametrize(
    "filename, payload, status",
    [
        ("spec.pdf", b"this is not a pdf at all", 422),
        ("spec.docx", b"this is not a docx at all", 422),
        ("old.doc", b"\xd0\xcf\x11\xe0legacy", 415),
        ("image.png", b"\x89PNG\r\n\x1a\n\x00\x00", 415),
        ("renamed.txt", b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR", 415),
        ("empty.txt", b"   \n", 422),
        ("empty.docx", None, 422),
    ],
)
def test_unreadable_documents_are_rejected(filename, payload, status):
    if payload is None:
        buf = io.BytesIO()
        Document().save(buf)
        payload = buf.getvalue()

    with pytest.raises(HTTPException) as excinfo:
        extract_document_text(filename, payload)

    assert excinfo.value.status_code == status
